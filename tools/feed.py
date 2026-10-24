#!/usr/bin/env python3
"""Đọc bảng hàng các sàn trên Google Drive → feed/<duan>.json (mã hóa AES-GCM bằng APP_KEY).
Cấu hình nguồn lấy từ biến môi trường SOURCES (JSON) hoặc file sources.json. Chạy: python feed.py [--plain]"""
import os,sys,re,io,csv,json,base64,unicodedata,urllib.request,urllib.parse,datetime,hashlib
import openpyxl
UA={'User-Agent':'Mozilla/5.0 feed-bot'}
LOG=[]
RELAY=os.environ.get('RELAY_URL',''); RTOK=os.environ.get('RELAY_TOKEN','')
def get(url,tries=3):
    if RELAY and ('docs.google.com' in url or 'drive.google.com' in url):
        st,b=_get(RELAY+('&' if '?' in RELAY else '?')+'t='+urllib.parse.quote(RTOK)+'&u='+urllib.parse.quote(url,safe=''),tries)
        if st==200:
            try:
                j=json.loads(b.decode('utf8'))
                if j.get('ok'): return int(j.get('code',200)),base64.b64decode(j['b64'])
                LOG.append(f"relay {j.get('code')} {url[:80]} {str(j.get('msg'))[:120]}")
                if j.get('code') in (401,403,404): return int(j['code']),b''
            except Exception as e: LOG.append(f'relay lỗi {e}')
        else: LOG.append(f'relay {st} {url[:80]}')
    return _get(url,tries)
def _get(url,tries=3):
    for i in range(tries):
        try:
            with urllib.request.urlopen(urllib.request.Request(url,headers=UA),timeout=60) as r: return r.status,r.read()
        except urllib.error.HTTPError as e:
            try: body=e.read()[:300]
            except Exception: body=b''
            LOG.append(f'{e.code} {url[:90]} {body[:160]!r}')
            if e.code in (429,500,502,503) and i<tries-1:
                import time; time.sleep(8*(i+1)); continue
            return e.code,b''
        except Exception as e:
            if i==tries-1: return 0,str(e).encode()
def norm(v): return re.sub(r'\s+',' ',str(v if v is not None else '')).strip()
def fold(s):
    s=unicodedata.normalize('NFD',norm(s).lower()); s=''.join(ch for ch in s if unicodedata.category(ch)!='Mn')
    return s.replace('đ','d')
def fillof(c):
    f=c.fill
    if not f or f.fill_type is None or f.fgColor is None: return None
    col=f.fgColor
    if col.type=='rgb' and col.rgb not in ('00000000','FFFFFFFF'): return col.rgb[-6:].upper()
    return None
DEFAULT_COLORS={'FF0000':'sold','FFFF00':'lock','00FF00':'check','999999':'check','B3E5A1':'lock'}
CODE_RX=re.compile(r'^([A-Z]{1,3}[0-9]{0,2}[A-Z]?[-.][0-9A-Z][0-9A-Z.\-]*|F2\d{4}|[A-ZĐ]{1,3}\d{1,3}-\d{1,3})$')
def num(s,ty=False):
    s=norm(s).replace(' ','')
    if not s or s=='-': return None
    if ty:
        s=s.replace(',','.') if s.count(',')==1 and s.count('.')==0 else s.replace(',','')
        try: return float(s)*1e9
        except: return None
    if re.match(r'^\d{1,3}(\.\d{3})+$',s) or re.match(r'^\d{1,3}(,\d{3})+$',s): return float(re.sub(r'[.,]','',s))
    s2=s.replace(',','.') if s.count(',')==1 and '.' not in s else s.replace(',','')
    try: return float(s2)
    except: return None
def area(s):
    s=norm(s)
    if re.match(r'^\d+,\d+$',s): s=s.replace(',','.')
    try: return round(float(s),2)
    except: return None
def stext(t):
    f=fold(t)
    if not f: return ''
    if re.search(r'da ban|sold|da coc|dat coc|da ky|ky hd|^coc\b',f): return 'sold'
    if re.search(r'\block|dang giu|^giu',f): return 'lock'
    if re.search(r'check|^thu\b|thu hoi',f): return 'check'
    return ''
HDR=[('code',r'^ma can( ho)?$|^can ho ten thuong mai|^ma can\b'),('typ',r'^loai( can| hinh| can ho| sp)?\b'),('tt',r'thong th|^dttt'),('tim',r'tim tuong'),
     ('land',r'dien tich ?dat|^dt dat'),('floor',r'san ?xd|dtxd|tong dien ?tich san'),('st',r'tinh trang'),('note',r'^ghi chu'),('coc',r'^dat coc$|^ngay coc'),('dl',r'^dai ly'),
     ('pre',r'truoc vat|chua (gom )?vat'),('vat',r'^vat$'),('kp',r'^kpbt$'),('dg',r'^don gia')]
def sheets_of(fid):
    st,b=get(f'https://docs.google.com/spreadsheets/d/{fid}/htmlview')
    if st!=200: return []
    t=b.decode('utf8','replace')
    return [(json.loads('"'+n+'"'),g) for n,g in re.findall(r'\{name: "((?:[^"\\]|\\.)*)", pageUrl: "[^"]*?gid=(\d+)',t)]
def parse_sheet(rows,ws,src):
    hi=None
    for i,r in enumerate(rows[:80]):
        if any(re.search(HDR[0][1],fold(c)) for c in r): hi=i; break
    if hi is None: return [],'không thấy dòng tiêu đề (cột Mã căn)'
    H=[fold(c) for c in rows[hi]]
    col={}
    for j,h in enumerate(H):
        if not h: continue
        for k,rx in HDR:
            if k not in col and re.search(rx,h): col[k]=j; break
    prx=src.get('price')
    if prx:
        for j,h in enumerate(H):
            if h and re.search(prx,h): col['g']=j; break
    leg={}
    colors=dict(DEFAULT_COLORS)
    fills={}
    if ws is not None:
        for r in ws.iter_rows():
            for c in r:
                v=norm(c.value); f=fillof(c)
                if not v: continue
                if c.row<=hi+1 and f and len(v)<26 and re.search(r'lock|ban|check|coc|giu',fold(v)): leg[f]=stext(v) if stext(v) else ('lock' if 'lock' in fold(v) else '')
                vv=v.upper().replace(' ','')
                if CODE_RX.match(vv) and vv not in fills: fills[vv]=f
        colors.update({k:v for k,v in leg.items() if v is not None})
        if leg: colors={k:v for k,v in colors.items() if k in leg or k not in DEFAULT_COLORS or DEFAULT_COLORS[k]=='sold'}
    out=[]
    for i in range(hi+1,len(rows)):
        r=rows[i]; cell=lambda k:(r[col[k]] if k in col and col[k]<len(r) else '')
        code=norm(cell('code')).upper().replace(' ','')
        if not CODE_RX.match(code): continue
        s=None
        for k in ('st','note'):
            if k in col and stext(cell(k)): s=stext(cell(k)); break
        f=fills.get(code)
        if s is None and f and colors.get(f): s=colors[f]
        if s is None and 'coc' in col and norm(cell('coc')): s='sold'
        if s is None and 'dl' in col and norm(cell('dl')): s='lock'
        rec={'c':code,'s':s or ''}
        g=num(cell('g'),src.get('ty')) if 'g' in col else None
        if g and g>1e8: rec['g']=round(g)
        for k in ('pre','vat','kp','dg'):
            if k in col:
                v=num(cell(k))
                if v and v>1e5: rec[k]=round(v)
        for k in ('typ','dl'):
            if norm(cell(k)): rec[k]=norm(cell(k))[:40]
        for k in ('tt','tim','land','floor'):
            if k in col:
                v=area(cell(k))
                if v and 5<v<3000: rec[k]=v
        out.append(rec)
    return out,''
def run_source(src,cache):
    fid=src['file']; res={'k':src['k'],'label':src['label'],'dv':src.get('dv',''),'ok':False,'n':0,'msg':''}
    if fid not in cache:
        wb=None
        for u in (f'https://docs.google.com/spreadsheets/d/{fid}/export?format=xlsx',f'https://drive.google.com/uc?export=download&id={fid}'):
            st,b=get(u)
            if st==200 and b[:2]==b'PK':
                try: wb=openpyxl.load_workbook(io.BytesIO(b),data_only=False); break
                except Exception as e: LOG.append(f'xlsx lỗi {e}')
            else: LOG.append(f'{st} {u[:80]} {b[:120]!r}')
        cache[fid]=wb
        cache[fid+'_sh']=sheets_of(fid)
    wb=cache[fid]; shs=cache[fid+'_sh']
    if wb is None: res['msg']='không tải được file (kiểm tra quyền chia sẻ)'; return res,[]
    want=[]
    for g in src.get('gids',[]):
        n=[n for n,gg in shs if gg==str(g)]
        if n: want.append((n[0],str(g)))
    for rx in src.get('sheets',[]):
        for n,g in shs:
            if re.search(rx,n) and (n,g) not in want: want.append((n,g))
    if not want: res['msg']='không thấy trang tính cần đọc'; return res,[]
    recs=[]; msgs=[]
    for n,g in want:
        st,b=get(f'https://docs.google.com/spreadsheets/d/{fid}/export?format=csv&gid={g}')
        if st!=200: msgs.append(f'{n}: lỗi tải {st}'); continue
        rows=list(csv.reader(io.StringIO(b.decode('utf8','replace'))))
        ws=wb[n] if n in wb.sheetnames else None
        rr,m=parse_sheet(rows,ws,src)
        if m: msgs.append(f'{n}: {m}')
        recs+=rr
    seen={}
    for r in recs: seen.setdefault(r['c'],r)
    recs=list(seen.values())
    for r in recs: r['k']=src['k']
    res.update(ok=bool(recs),n=len(recs),msg='; '.join(msgs),sheets=[n for n,_ in want])
    return res,recs

# ===== SalePro (API công khai, không cần đăng nhập): quỹ đang mở bán theo dự án =====
SALEPRO={'pb':(1402,''),'ab':(1304,''),'ih':(1312,''),'sc':(1264,'SUN CENTRO TOWN'),'sb':(1361,'*'),'hlx':(1334,'*'),'ho':(1343,'*'),
  'dk':(594,'*'),'xi':(1322,''),'js':(1300,''),'ri':(181,'*'),'ga':(1324,'*'),'am':(1329,''),'fc':(1397,''),'ru':(1439,'*'),'cv':(1462,''),'np':(1408,'*')}
SPB='https://salepro.com/api/rest_api/pu/property/master/data/search_tindang_data/'
SPL='https://salepro.com/api/rest_api/pu/property/master/data/search_tindang_low_rise_data/'
SPUA={'User-Agent':'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/130.0 Safari/537.36','Accept':'application/json, text/plain, */*','Referer':'https://salepro.com/'}
TYP={'STUDIO':'ST','STUDIO+':'ST+','1 PHÒNG NGỦ':'1BR','1 PHÒNG NGỦ +':'1BR+','2 PHÒNG NGỦ':'2BR','2 PHÒNG NGỦ +':'2BR+','3 PHÒNG NGỦ':'3BR','3 PHÒNG NGỦ +':'3BR+','4 PHÒNG NGỦ':'4BR'}
QN={'QC':'Quỹ chéo'}
def _spget(q,low=False):
    for i in range(3):
        try:
            with urllib.request.urlopen(urllib.request.Request((SPL if low else SPB)+'?'+q,headers=SPUA),timeout=60) as r: return json.loads(r.read().decode('utf8')).get('data') or []
        except Exception as e:
            err=e
            import time; time.sleep(3*(i+1))
    raise err
def salepro(pid,sub):
    import math
    low=sub=='*'; sub='' if low else sub
    d=_spget(f'project_id={pid}&page=1',low); n=int(d[0]['total_records']) if d else 0
    rows=list(d)
    import concurrent.futures as cf
    with cf.ThreadPoolExecutor(6) as ex:
        for part in ex.map(lambda p:_spget(f'project_id={pid}&page={p}',low),range(2,math.ceil(n/24)+1)): rows+=part
    codes={r['product_code'] for r in rows}
    if n and len(codes)<n*0.9: raise Exception(f'thiếu dữ liệu {len(codes)}/{n}')
    out=[];seen=set()
    for r in rows:
        if sub and norm(r.get('subdivision_name')).upper()!=sub: continue
        st_='' if (r.get('state_line')=='available' and not r.get('is_not_show') and r.get('state')=='public') else ('sold' if (r.get('state_line')=='sold' or 'bán' in norm(r.get('state_line_name')).lower()) else 'off')
        c=norm(r['product_code'])
        if c in seen: continue
        seen.add(c)
        num_=lambda v:(round(float(v),2) if v not in (None,'') and float(v)>0 else None)
        pr=lambda v:(int(round(float(v))) if v not in (None,'') and float(v)>1000 else None)
        tw=((lambda c:(c.split('.')[0] if '.' in c else c.rsplit('-',1)[0]))(norm(r['product_code'])) if low else re.sub(r'^To[àa]\s*','',norm(r.get('building_name')),flags=re.I)) or None
        out.append([c,tw,norm(r.get('number_of_floor')) or None,norm(r.get('apartment_axis_code')) or None,
            TYP.get(norm(r.get('apartment_type_name')).upper(),norm(r.get('apartment_type_name'))),
            num_(r.get('carpet_area')),num_(r.get('built_up_area')),norm(r.get('base_direction_name')),'',
            pr(r.get('apartment_price')),norm(r.get('price_description')),QN.get(norm(r.get('fund_groups_sell_code')),norm(r.get('fund_groups_sell_code'))),pr(r.get('early_payment_price')),st_])
    out.sort(key=lambda x:x[13]!='')
    return {'n':n,'u':out}
def sp_merge(new,old,day):
    """Không xóa căn: căn đã từng có mà nay không còn trên SalePro -> 'off' (chưa mở bán / rút khỏi quỹ); ghi ngày đổi tình trạng."""
    om={r[0]:r for r in (old or [])}; out=[]; seen=set()
    for r in new:
        st=r[13]; p=om.get(r[0]); od=(p[14] if p and len(p)>14 else '')
        ost=(p[13] if p and len(p)>13 else '') if p else None
        r=r[:14]+[od if (p is not None and ost==st) else (day if st else '')]
        out.append(r); seen.add(r[0])
    for c,p in om.items():
        if c in seen: continue
        p=list(p)+['']*(15-len(p))
        if p[13]=='' : p[13]='off'; p[14]=day
        out.append(p[:15])
    return out
def enc(obj,key):
    from cryptography.hazmat.primitives.ciphers.aead import AESGCM
    iv=os.urandom(12); ct=AESGCM(base64.b64decode(key)).encrypt(iv,json.dumps(obj,ensure_ascii=False,separators=(',',':')).encode(),None)
    return base64.b64encode(iv+ct).decode()

# ===== SalePro công khai: nội dung các tab dự án (tổng quan, vị trí, 360, phân khu, mặt bằng, tiện ích...) – không lấy tin tức =====
from html.parser import HTMLParser
_OKT={'p','br','b','strong','i','em','u','ul','ol','li','h2','h3','h4','h5','img','a','table','thead','tbody','tr','td','th','span','div','blockquote','figure','figcaption'}
_PH=re.compile(r'(?<!\d)(?:\+?84|0)(?:[\s.\-]?\d){8,10}(?!\d)')
class _San(HTMLParser):
    def __init__(s): super().__init__(convert_charrefs=True); s.o=[]; s.skip=0
    def handle_starttag(s,t,a):
        if t in('script','style','iframe','object','embed','form','svg'): s.skip+=1; return
        if s.skip or t not in _OKT: return
        a=dict(a); keep=''
        if t=='img' and str(a.get('src','')).startswith('https://'): keep=f' src="{a["src"]}" loading="lazy" style="max-width:100%"'
        elif t=='img': return
        if t=='a' and str(a.get('href','')).startswith('https://'): keep=f' href="{a["href"]}" target="_blank" rel="noopener"'
        if t in('td','th') and a.get('colspan','').isdigit(): keep=f' colspan="{a["colspan"]}"'
        s.o.append(f'<{t}{keep}>')
    def handle_endtag(s,t):
        if t in('script','style','iframe','object','embed','form','svg'): s.skip=max(0,s.skip-1); return
        if not s.skip and t in _OKT and t not in('img','br'): s.o.append(f'</{t}>')
    def handle_data(s,d):
        if not s.skip: s.o.append(_PH.sub('',d).replace('&','&amp;').replace('<','&lt;').replace('>','&gt;'))
def san(h):
    if not h or not isinstance(h,str): return ''
    x=_San(); x.feed(h); x.close(); r=''.join(x.o)
    r=re.sub(r'<(p|li|span|div|h\d)>\s*(&nbsp;|\s)*</\1>','',r)
    return r.strip()
SPM='https://salepro.com/api/rest_api/pu/property/master/data/'
def _spj(path,tries=3):
    err=None
    for i in range(tries):
        try:
            with urllib.request.urlopen(urllib.request.Request(path if path.startswith('http') else SPM+path,headers=SPUA),timeout=60) as r: return json.loads(r.read().decode('utf8'))
        except Exception as e:
            err=e; import time; time.sleep(3*(i+1))
    raise err
def _img(o): return (o or {}).get('image_url') if isinstance(o,dict) else None
def _yt(u):
    m=re.search(r'(?:youtu\.be/|v=|embed/|shorts/)([\w-]{11})',u or ''); return m.group(1) if m else None
def landing(pid):
    D=_spj(f'get_project_detail/?project_id={pid}').get('data') or {}
    I=D.get('project_info') or {}
    if not I: raise Exception('không có project_info')
    slug=I.get('project_slug')
    try: L=(_spj(f'get_project_landing/?project_id={slug}').get('data') or [{}])[0] or {}
    except Exception: L={}
    low=I.get('project_type')!='high_rise'
    try: TN=_spj(f'get_toanha_data/?project_id={pid}').get('data') or []
    except Exception: TN=[]
    try: PK=_spj(f'get_phankhu_data/?project_id={pid}').get('data') or []
    except Exception: PK=[]
    f=lambda k:norm(I.get(k)) or None
    loc=re.search(r'src="(https://www\.google\.com/maps/embed[^"]+)"',I.get('iframe_location') or '')
    o={'name':f('project_name'),'type':'Thấp tầng' if low else 'Cao tầng',
       'facts':[[a,v] for a,v in [('Pháp lý / hình thức sở hữu',f('ownership_type')),('Tổng diện tích',f('total_area')),('Mật độ xây dựng',f('construction_density')),('Quy mô',f('development_scale')),('Chủ đầu tư',f('investor')),('Khởi công',f('build_date')),('Tiến độ',f('progress'))] if v]
             +[[norm(x.get('name')),norm(x.get('value'))] for x in (L.get('project_info') or []) if x.get('name') and x.get('value')],
       'desc':san(L.get('project_html_description') or I.get('short_description') or ''),'ovh':san(I.get('project_overview') or I.get('project_detail') or ''),
       'img':_img(L.get('project_image')),'ban':[x.get('image_url') for x in (L.get('banner_desktop') or []) if x.get('image_url')][:4],
       'loc':san(I.get('location') or ''),'map':loc.group(1) if loc else None,
       'v360':I.get('image_360') if str(I.get('image_360') or '').startswith('https://') else None,
       'i360':[[m.get('image_name') or x.get('image_coordinate_name') or '',m['image_url']] for x in (D.get('img_360_data') or []) for m in (x.get('image_data') or []) if m.get('image_url')][:30],
       'fp':[[norm(x.get('floorplan_name')),x['image_url']] for x in (L.get('floor_plan') or []) if x.get('image_url')],
       'pos':[[norm(x.get('image_coordinate_name')),m['image_url']] for x in (D.get('image_data') or []) for m in (x.get('image_data') or []) if m.get('image_url')][:40],
       'prod':[[norm(x.get('product_name')),norm(x.get('product_price')),san(x.get('product_features') or '') if isinstance(x.get('product_features'),str) else '',x.get('image_url')] for x in (L.get('product') or [])],
       'am':[[norm(x.get('amenity_name')),x.get('image_url')] for x in (L.get('amenities') or [])],
       'vid':[[norm(L.get(p+'_title')),norm(L.get(p+'_content')),_yt(L.get(p+'_video_url'))] for p in ('discovery','entertainment') if L.get(p+'_title') or _yt(L.get(p+'_video_url'))],
       'csi':[u for u in (I.get('sale_policy_images') or []) if isinstance(u,str) and u.startswith('https://')],
       'lk':[[norm(x.get('link_utility_name')),x['link_utility_url']] for x in (D.get('link_general') or [])+(D.get('link_utility') or []) if str(x.get('link_utility_url','')).startswith('https://')],
       'tn':[[norm(x.get('building_name')),norm(x.get('subdivision_name')),norm(x.get('building_total_floor')),x.get('total_apartment'),norm(x.get('handover_standard')),_img(x.get('avatar'))] for x in TN if x.get('state','public')=='public'],
       'pk':[[norm(x.get('subdivision_name')),norm(x.get('total_area')),norm(x.get('total_apartment')),norm(x.get('style_build')),_img(x.get('avatar'))] for x in PK if x.get('state','public')=='public']}
    o['desc']=_PH.sub('',o['desc']); seen=set(); o['facts']=[x for x in o['facts'] if not (x[0].lower() in seen or seen.add(x[0].lower()))]
    return o
def run_landing(out_dir,key,now,plain=False):
    st={}
    for duan,(pid,sub) in SALEPRO.items():
        p=os.path.join(out_dir,'land_'+duan+'.json')
        try: o=landing(pid)
        except Exception as e: st[duan]='lỗi '+str(e)[:80]; continue
        h=hashlib.sha256(json.dumps(o,sort_keys=True,ensure_ascii=False).encode()).hexdigest()[:16]
        try: old=json.load(open(p))
        except Exception: old={}
        if old.get('h')==h and not plain: st[duan]='không đổi'; continue
        body={'at':now.strftime('%H:%M %d/%m/%Y'),'o':o}
        json.dump({'h':h,'at':body['at'],'d':body} if plain or not key else {'h':h,'at':body['at'],'e':enc(body,key)},open(p,'w',encoding='utf8'),ensure_ascii=False,separators=(',',':'))
        st[duan]='cập nhật'
    print('LANDING',json.dumps(st,ensure_ascii=False),file=sys.stderr)
def main():
    plain='--plain' in sys.argv
    key=os.environ.get('APP_KEY','')
    here=os.path.dirname(os.path.abspath(__file__)); encp=os.path.join(here,'sources.enc')
    if key and os.path.exists(encp):
        from cryptography.hazmat.primitives.ciphers.aead import AESGCM
        raw=base64.b64decode(open(encp).read().strip()); cfg=json.loads(AESGCM(base64.b64decode(key)).decrypt(raw[:12],raw[12:],None))
    elif os.environ.get('SOURCES'): cfg=json.loads(os.environ['SOURCES'])
    else: cfg=json.load(open(os.path.join(here,'sources.json'),encoding='utf8'))
    out_dir=os.environ.get('OUT','feed'); os.makedirs(out_dir,exist_ok=True)
    now=datetime.datetime.now(datetime.timezone(datetime.timedelta(hours=7)))
    def old_feed(duan):
        try:
            o=json.load(open(os.path.join(out_dir,duan+'.json'),encoding='utf8'))
            if 'd' in o: return o['d']
            if key and 'e' in o:
                from cryptography.hazmat.primitives.ciphers.aead import AESGCM
                raw=base64.b64decode(o['e']); return json.loads(AESGCM(base64.b64decode(key)).decrypt(raw[:12],raw[12:],None))
        except Exception: return None
    olds={}
    cache={}; per={}
    for src in cfg:
        res,recs=run_source(src,cache)
        if not res['ok']:
            od=olds.setdefault(src['duan'],old_feed(src['duan']))
            if od:
                os_=next((x for x in od.get('src',[]) if x['k']==src['k'] and x.get('ok')),None)
                orec=[r for r in od.get('u',[]) if r.get('k')==src['k']]
                if os_ and orec:
                    print(f"{src['k']:12} lỗi lần này ({res['msg']}) – giữ dữ liệu lần trước {os_.get('at')}",file=sys.stderr)
                    res=dict(os_); res['stale']=1; res['msg']='Lần đọc gần nhất lỗi, đang dùng dữ liệu lúc '+str(os_.get('at')); recs=orec
                    per.setdefault(src['duan'],{'src':[],'u':[]})['src'].append(res); per[src['duan']]['u']+=recs; continue
        res['at']=now.strftime('%H:%M %d/%m/%Y')
        if res['ok'] and src.get('full',True):
            od=olds.setdefault(src['duan'],old_feed(src['duan']))
            have={r['c'] for r in recs}; kept=0
            for r in (od or {}).get('u',[]):
                if r.get('k')!=src['k'] or r['c'] in have: continue
                r=dict(r)
                if not r.get('x'): r['x']=now.strftime('%d/%m/%Y')
                r['s']='sold'; recs.append(r); have.add(r['c']); kept+=1
            if kept: res['msg']=(res['msg']+'; ' if res['msg'] else '')+f'{kept} căn không còn trong bảng (giữ lại, ghi đã bán)'
        print(f"{src['k']:12} {res['n']:4} {res['msg']}",file=sys.stderr)
        d=per.setdefault(src['duan'],{'src':[],'u':[]})
        d['src'].append(res); d['u']+=recs
    for duan,(pid,sub) in SALEPRO.items():
        d=per.setdefault(duan,{'src':[],'u':[]})
        try:
            r=salepro(pid,sub)
            if not any(x[13]=='' for x in r['u']): raise Exception('không có căn đang mở bán')
            od=olds.setdefault(duan,old_feed(duan)); osp=(od or {}).get('sp') or {}
            u=sp_merge(r['u'],osp.get('u'),now.strftime('%d/%m/%Y'))
            d['sp']={'date':now.strftime('%d/%m/%Y'),'at':now.strftime('%H:%M %d/%m/%Y'),'u':u}; d['spnew']=1
            print(f"salepro_{duan:5} {sum(1 for x in u if x[13]==''):4} căn đang mở bán / {len(u)} căn lưu",file=sys.stderr)
        except Exception as e:
            od=olds.setdefault(duan,old_feed(duan))
            if od and od.get('sp'): d['sp']=od['sp']
            print(f"salepro_{duan:5} lỗi {str(e)[:160]} – {'giữ bản cũ' if d.get('sp') else 'chưa có dữ liệu'}",file=sys.stderr)
    status={}
    for duan,d in per.items():
        if not any(x.get('ok') for x in d['src']) and not d.get('spnew') and os.path.exists(os.path.join(out_dir,duan+'.json')):
            status[duan]='lỗi đọc – giữ file cũ'; continue
        body={'v':1,'at':now.strftime('%H:%M %d/%m/%Y'),'src':d['src'],'u':d['u']}
        if d.get('sp'): body['sp']=d['sp']
        h=hashlib.sha256(json.dumps({'src':[{k:v for k,v in s.items() if k!='at'} for s in d['src']],'u':d['u'],'sp':(d.get('sp') or {}).get('u')},sort_keys=True,ensure_ascii=False).encode()).hexdigest()[:16]
        body['h']=h
        p=os.path.join(out_dir,duan+'.json')
        try: old=json.load(open(p))
        except Exception: old={}
        if old.get('h')==h and not plain: status[duan]='không đổi'; continue
        data={'h':h,'at':body['at'],'d':body} if plain or not key else {'h':h,'at':body['at'],'e':enc(body,key)}
        json.dump(data,open(p,'w',encoding='utf8'),ensure_ascii=False,separators=(',',':'))
        status[duan]='cập nhật'
    print(json.dumps(status,ensure_ascii=False),file=sys.stderr)
    try: run_landing(out_dir,key,now,plain)
    except Exception as e: print('LANDING lỗi',e,file=sys.stderr)
    for l in LOG[:40]: print('LOG',l,file=sys.stderr)
if __name__=='__main__': main()
